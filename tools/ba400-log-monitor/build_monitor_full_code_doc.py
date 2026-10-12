from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/ricardomontanezmiranda/Desktop/Biosystems Project")
MONITOR_DIR = ROOT / "tools" / "ba400-log-monitor"
OUTPUT_DIR = ROOT / "outputs" / "monitor-docs"
OUTPUT_PATH = OUTPUT_DIR / "Codigo_Fuente_Completo_Documentado_AX00_Log_Monitor.docx"

FILES = {
    "full-monitor.mjs": MONITOR_DIR / "full-monitor.mjs",
    "monitor.mjs": MONITOR_DIR / "monitor.mjs",
    "consumption-monitor.mjs": MONITOR_DIR / "consumption-monitor.mjs",
}

TITLE_BLUE = RGBColor(0x2E, 0x74, 0xB5)
TITLE_DARK = RGBColor(0x1F, 0x4D, 0x78)
TEXT_GRAY = RGBColor(0x5A, 0x5A, 0x5A)
LINE_GRAY = RGBColor(0x88, 0x88, 0x88)
LIGHT_FILL = "EAF1F8"
CODE_FILL = "F7F7F7"


FILE_NOTES = {
    "full-monitor.mjs": {
        "purpose": "Supervisor del sistema. Arranca, vigila y reinicia los monitores de errores y consumos.",
        "why": "Este archivo existe para desacoplar los dos monitores productivos y dejar que el servicio se recupere solo cuando un proceso cae.",
        "hotspots": [
            "normalizeConfig() convierte el JSON en rutas ejecutables y resuelve placeholders.",
            "pipeWithPrefix() preserva el origen de cada linea de log al mezclar stdout y stderr.",
            "MonitorSupervisor.startMonitor() define la politica real de reinicio automatico.",
            "main() instala el apagado limpio ante SIGINT y SIGTERM.",
        ],
    },
    "monitor.mjs": {
        "purpose": "Monitor de errores en tiempo real sobre archivos TraceComm*.LAx00.",
        "why": "Lee solo el crecimiento nuevo del log, deduce el numero de serie real del equipo y publica tanto historico como estado actual de error.",
        "hotspots": [
            "resolveSourceDirectory() y rankErrorDirectoryCandidates() deciden cual carpeta es la instalacion activa.",
            "readIncrementalChunk() y pending_fragment permiten seguir un archivo abierto sin releerlo completo.",
            "parseTraceCommLine() concentra la semantica del protocolo TraceComm.",
            "buildEvent(), buildCurrentErrorStateRow() y upload*() definen el contrato con Supabase.",
        ],
    },
    "consumption-monitor.mjs": {
        "purpose": "Monitor de consumos mensuales desde archivos CSV de LogConsum.",
        "why": "Recalcula totales absolutos por mes y equipo para reactivos, rotores y estado operativo, evitando tablas inmanejables en Supabase.",
        "hotspots": [
            "resolveCandidateFiles() y shouldReprocessFile() gobiernan que archivos se reevalúan.",
            "processReagentFile() resume cada pipeteo en agregados mensuales por prueba.",
            "processRotorFile() resume los cambios de rotor sin subir evento por evento.",
            "buildPipelineSignature() permite invalidar el estado local cuando cambia el contrato de salida.",
        ],
    },
}

FUNCTION_NOTES = {
    "full-monitor.mjs": {
        "parseArgs": "Lee `--config` y `--once` para controlar el modo de arranque.",
        "readJson": "Carga el archivo JSON de configuracion.",
        "timestamp": "Genera la marca de tiempo uniforme para logging.",
        "log": "Emite mensajes del supervisor.",
        "sleep": "Pausa asincrona usada antes de reinicios.",
        "expandPlaceholders": "Resuelve placeholders de rutas y variables de entorno.",
        "deepExpandPlaceholders": "Aplica la expansion sobre objetos y arreglos anidados.",
        "normalizeConfig": "Convierte el JSON crudo en rutas y parametros ejecutables.",
        "ensureLogDir": "Crea la carpeta de logs si no existe.",
        "pipeWithPrefix": "Prefija cada linea para distinguir de que monitor proviene.",
        "MonitorSupervisor": "Clase responsable de crear, vigilar, reiniciar y detener procesos hijos.",
        "main": "Punto de entrada que inicializa el supervisor y el apagado limpio.",
    },
    "monitor.mjs": {
        "parseArgs": "Lee los argumentos de linea de comandos.",
        "readJson": "Carga configuracion o catalogos auxiliares.",
        "writeJson": "Persiste el estado local del monitor.",
        "timestamp": "Genera timestamps de log.",
        "log": "Centraliza la salida de diagnostico.",
        "toLocalDateKey": "Convierte fechas a `YYYYMMDD` para priorizar el TraceComm del dia.",
        "sleep": "Pausa breve para reintentos.",
        "expandPlaceholders": "Expande placeholders de configuracion.",
        "deepExpandPlaceholders": "Aplica esa expansion sobre estructuras anidadas.",
        "getValueByPath": "Extrae valores por dot-path para el mapeo de salida.",
        "removeUndefinedProperties": "Elimina campos `undefined` antes de subir filas.",
        "inferEquipmentModelFromSerial": "Infiere BA400, BA200, A15 o A25 segun prefijo de serie.",
        "extractPreferredTraceCommSerial": "Busca ASN/SN y prioriza ASN como identidad efectiva.",
        "normalizeCatalogErrorType": "Normaliza la severidad del catalogo a `ok`, `warning` o `fatal`.",
        "resolveMessageTypeFromCatalog": "Escoge la severidad consolidada cuando hay varios codigos.",
        "fileExists": "Comprueba si existe el archivo de estado.",
        "loadState": "Carga o inicializa el estado local.",
        "resolveSourceDirectory": "Encuentra la carpeta de log realmente activa.",
        "extractTraceDateKeyFromName": "Saca la fecha desde el nombre del archivo TraceComm.",
        "rankErrorDirectoryCandidates": "Ordena carpetas candidatas segun fecha vigente y actividad reciente.",
        "resolveActiveFile": "Escoge el archivo TraceComm activo dentro de la carpeta elegida.",
        "getOrCreateFileState": "Crea el bloque de estado por archivo si aun no existe.",
        "readIncrementalChunk": "Lee solo el rango de bytes nuevo desde el ultimo offset.",
        "readTextSlice": "Convierte un segmento de bytes a texto con el encoding configurado.",
        "extractTraceCommIdentity": "Extrae ASN/SN desde un bloque de texto.",
        "bootstrapTraceCommIdentity": "Aprende la serie desde cabecera o cola aunque no haya errores nuevos.",
        "parseErrorCodes": "Detecta codigos `E:(codigo)` en una linea.",
        "parseTraceCommLine": "Interpreta una linea TraceComm y decide si hay evento, limpieza o permanencia.",
        "buildMatchedDescriptions": "Cruza los codigos con el catalogo local.",
        "resolveCurrentStateStatus": "Calcula el estado final actual del equipo.",
        "buildEvent": "Construye la fila historica para `monitoreo_errores_equipos`.",
        "buildCurrentErrorStateRow": "Construye la fila consolidada para `estado_errores_equipo_actual`.",
        "mapEventToRow": "Adapta el payload interno al esquema SQL configurado.",
        "uploadBatch": "Inserta filas historicas por lotes.",
        "upsertBatch": "Hace upsert sobre tablas de estado actual.",
        "uploadEvents": "Parte eventos en lotes y los inserta.",
        "uploadCurrentStates": "Deduplica por serie y sincroniza el estado actual.",
        "pollOnce": "Ejecuta un ciclo completo de lectura, parseo y sincronizacion.",
        "normalizeConfig": "Normaliza la configuracion completa del monitor.",
        "main": "Bucle principal con manejo de errores y frecuencia de sondeo.",
    },
    "consumption-monitor.mjs": {
        "parseArgs": "Lee `--config` y `--once`.",
        "readJson": "Carga configuraciones y estado.",
        "writeJson": "Guarda el estado local de consumos.",
        "fileExists": "Comprueba la existencia del archivo de estado.",
        "timestamp": "Genera timestamps de log.",
        "log": "Escribe mensajes de diagnostico.",
        "sleep": "Pausa breve en reintentos.",
        "expandPlaceholders": "Expande placeholders de configuracion.",
        "deepExpandPlaceholders": "Hace la expansion profunda de placeholders.",
        "inferEquipmentModelFromSerial": "Deduce el modelo a partir del prefijo de serie.",
        "normalizeText": "Limpia texto crudo.",
        "normalizeCsvCell": "Limpia celdas CSV.",
        "parseNumber": "Convierte numeros de texto a `number`.",
        "parseBoolean": "Normaliza booleanos del CSV.",
        "parseLocalTimestamp": "Interpreta timestamps locales de consumo.",
        "sha1": "Genera hashes estables para llaves tecnicas.",
        "buildPipelineSignature": "Firma el contrato del pipeline para invalidar estado si cambia.",
        "parseCsvLine": "Implementa un parser CSV ligero con comillas.",
        "mapCsvRow": "Mapea encabezados a valores.",
        "normalizeSampleClassKey": "Reduce `SampleClass` a una clave canonica.",
        "toLocalMonthKey": "Genera la llave `YYYYMM` del mes local.",
        "readTextFileWithRetries": "Lee archivos bloqueados con reintentos tolerantes.",
        "parseConsumptionFileName": "Extrae serie, tipo y mes del nombre del archivo.",
        "extractConsumptionMonthKey": "Obtiene la llave mensual desde el nombre.",
        "extractConsumptionKind": "Clasifica archivo de reactivo o rotor.",
        "rankConsumptionDirectoryCandidates": "Prioriza carpetas con archivos vigentes del mes actual.",
        "resolveSourceDirectory": "Encuentra la carpeta LogConsum realmente activa.",
        "resolveCandidateFiles": "Lista los archivos fuente validos que pueden reprocesarse.",
        "loadState": "Carga el estado local del monitor de consumos.",
        "getOrCreateFileState": "Inicializa el estado asociado a un archivo mensual.",
        "shouldReprocessFile": "Decide si un archivo requiere nueva lectura completa.",
        "ensureEquipmentState": "Inicializa el acumulador de estado operativo por equipo.",
        "updateEquipmentStateFromRow": "Actualiza pack ISE y electrodos con la fila mas reciente.",
        "buildEquipmentStateRow": "Materializa el estado operativo final para Supabase.",
        "createReagentSummary": "Crea el acumulador mensual por prueba.",
        "updateReagentSummary": "Incrementa conteos y volúmenes del resumen de reactivos.",
        "finalizeReagentSummary": "Cierra el resumen con redondeos y metadatos finales.",
        "createRotorSummary": "Crea el acumulador mensual de rotores.",
        "updateRotorSummary": "Actualiza conteo y fechas de cambios de rotor.",
        "finalizeRotorSummary": "Materializa la fila final de rotor.",
        "resolveEffectiveIdentity": "Escoge la serie efectiva desde CSV, configuracion o nombre.",
        "mergeEquipmentStateMaps": "Fusiona estados parciales conservando el dato mas reciente.",
        "processReagentFile": "Lee un CSV de reactivos y produce acumulados mensuales.",
        "processRotorFile": "Lee un CSV de rotores y produce acumulados de cambios.",
        "processFile": "Despacha al parser correcto segun el tipo de archivo.",
        "chunkRows": "Parte lotes grandes para upsert.",
        "upsertRows": "Hace upsert de filas en Supabase.",
        "normalizeConfig": "Normaliza la configuracion del monitor de consumos.",
        "pollOnce": "Ejecuta un ciclo completo de agregacion y sincronizacion.",
        "main": "Bucle principal con la frecuencia configurada.",
    },
}


@dataclass
class Definition:
    kind: str
    name: str
    line: int


@dataclass
class Chunk:
    title: str
    line_start: int
    line_end: int
    description: str
    lines: list[str]


def set_run_font(run, *, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_fixed_layout(table, widths_in_inches, *, header_fill=LIGHT_FILL):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    width_dxa_values = [int(width * 1440) for width in widths_in_inches]
    total_dxa = sum(width_dxa_values)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width_dxa in width_dxa_values:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width_dxa))
        tbl_grid.append(col)

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths_in_inches[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa_values[col_index]))
            tc_w.set(qn("w:type"), "dxa")
            if row_index == 0 and header_fill:
                shade_cell(cell, header_fill)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, name="Calibri", size={1: 16, 2: 13, 3: 11.5}[level], color=TITLE_BLUE, bold=True)
    if level == 1:
        set_paragraph_spacing(paragraph, before=14, after=7, line=1.0)
    elif level == 2:
        set_paragraph_spacing(paragraph, before=10, after=5, line=1.0)
    else:
        set_paragraph_spacing(paragraph, before=6, after=4, line=1.0)
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=6, line=1.05)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=10.5, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, size=10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(paragraph, before=0, after=3, line=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        paragraph = cell.paragraphs[0]
        set_paragraph_spacing(paragraph, before=0, after=0, line=1.0)
        run = paragraph.add_run(header)
        set_run_font(run, size=9.5, bold=True, color=TITLE_DARK)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            set_paragraph_spacing(paragraph, before=0, after=0, line=1.0)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.25)
    set_table_fixed_layout(table, widths)
    doc.add_paragraph()
    return table


def set_page_layout(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)


def set_document_defaults(doc):
    set_page_layout(doc)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)

    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.color.rgb = TITLE_BLUE
        style.font.bold = True

    header = doc.sections[0].header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(header_paragraph, before=0, after=0, line=1.0)
    header_run = header_paragraph.add_run("AX00 Log Monitor | Codigo fuente completo documentado")
    set_run_font(header_run, size=8.5, color=TEXT_GRAY)

    footer = doc.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer_paragraph, before=0, after=0, line=1.0)
    footer_run = footer_paragraph.add_run("Generado el 20/08/2026 desde el codigo fuente del proyecto")
    set_run_font(footer_run, size=8.5, color=TEXT_GRAY)


def parse_definitions(source_text: str) -> list[Definition]:
    definitions: list[Definition] = []
    for line_number, line in enumerate(source_text.splitlines(), start=1):
        match = re.match(r"^(function|async function|class)\s+([A-Za-z0-9_]+)", line)
        if match:
            definitions.append(Definition(kind=match.group(1), name=match.group(2), line=line_number))
    return definitions


def chunk_description(file_name: str, chunk_name: str) -> str:
    if chunk_name == "__preamble__":
        return "Bloque inicial con imports, constantes, regex y configuracion base necesaria para que el archivo funcione."
    if chunk_name == "__epilogue__":
        return "Bloque final con invocaciones, exportaciones o cierre del script."
    return FUNCTION_NOTES.get(file_name, {}).get(
        chunk_name,
        "Bloque funcional del archivo. Conviene leerlo junto con el flujo general descrito al inicio de esta seccion.",
    )


def chunk_title(definition: Definition | None) -> str:
    if definition is None:
        return "Bloque inicial"
    if definition.kind == "class":
        return f"class {definition.name}"
    return f"{definition.name}()"


def build_chunks(file_name: str, source_text: str) -> list[Chunk]:
    lines = source_text.splitlines()
    definitions = parse_definitions(source_text)
    chunks: list[Chunk] = []

    if definitions and definitions[0].line > 1:
        preamble_end = definitions[0].line - 1
        chunks.append(
            Chunk(
                title="Bloque inicial",
                line_start=1,
                line_end=preamble_end,
                description=chunk_description(file_name, "__preamble__"),
                lines=lines[0:preamble_end],
            )
        )

    for index, definition in enumerate(definitions):
        start = definition.line
        end = definitions[index + 1].line - 1 if index + 1 < len(definitions) else len(lines)
        chunks.append(
            Chunk(
                title=chunk_title(definition),
                line_start=start,
                line_end=end,
                description=chunk_description(file_name, definition.name),
                lines=lines[start - 1:end],
            )
        )

    if not definitions:
        chunks.append(
            Chunk(
                title="Archivo completo",
                line_start=1,
                line_end=len(lines),
                description="Archivo sin definiciones detectables por regex simple.",
                lines=lines,
            )
        )

    return chunks


def add_cover(doc, summaries):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title, before=18, after=4, line=1.0)
    run = title.add_run("Codigo fuente completo documentado")
    set_run_font(run, size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(subtitle, before=0, after=14, line=1.0)
    run = subtitle.add_run("full-monitor.mjs, monitor.mjs y consumption-monitor.mjs con explicacion integrada")
    set_run_font(run, size=13, color=TEXT_GRAY)

    add_body(doc, "Proyecto: AX00 Log Monitor para monitoreo nacional de errores y consumos.", bold_lead="Proyecto:")
    add_body(doc, "Fecha: 20 de agosto de 2026.", bold_lead="Fecha:")
    add_body(doc, "Formato: documento de referencia con descripcion por archivo, mapa de bloques y listing completo del codigo fuente.", bold_lead="Formato:")
    add_body(doc, "Archivos incluidos: " + ", ".join(summary["file_name"] for summary in summaries), bold_lead="Archivos incluidos:")

    doc.add_paragraph()
    add_body(
        doc,
        "Esta version no resume el codigo: lo reproduce completo dentro del Word. Cada archivo se abre con una guia de lectura, luego se presenta un mapa de bloques con rangos de linea y finalmente el listing integro numerado para facilitar soporte, auditoria y mantenimiento.",
    )


def add_reading_guide(doc):
    add_heading(doc, "1. Como leer este documento", 1)
    add_bullet(doc, "Cada archivo empieza con su objetivo operativo y los puntos de mantenimiento mas sensibles.")
    add_bullet(doc, "El mapa de bloques indica en que lineas empieza y termina cada funcion o clase.")
    add_bullet(doc, "El codigo aparece completo, con numeracion de linea para poder referenciarlo en soporte o revisiones.")
    add_bullet(doc, "Las notas descriptivas no alteran el source real: sirven para explicar la intencion de cada bloque.")

    add_heading(doc, "1.1 Archivos cubiertos", 2)
    add_table(
        doc,
        ["Archivo", "Lineas", "Responsabilidad principal"],
        [[summary["file_name"], summary["line_count"], FILE_NOTES[summary["file_name"]]["purpose"]] for summary in doc._monitor_summaries],  # type: ignore[attr-defined]
        [2.2, 0.8, 5.0],
    )


def add_code_line(doc, line_number: int, code: str):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=0, line=0.95)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), CODE_FILL)
    shd.set(qn("w:val"), "clear")

    number_run = paragraph.add_run(f"{line_number:>4} | ")
    set_run_font(number_run, name="Courier New", size=7.5, color=LINE_GRAY)
    code_run = paragraph.add_run(code.replace("\t", "  "))
    set_run_font(code_run, name="Courier New", size=7.8, color=RGBColor(0, 0, 0))


def summarize_file(file_name: str, source_path: Path):
    source_text = source_path.read_text(encoding="utf8")
    definitions = parse_definitions(source_text)
    return {
        "file_name": file_name,
        "path": str(source_path),
        "line_count": len(source_text.splitlines()),
        "definition_count": len(definitions),
        "chunks": build_chunks(file_name, source_text),
        "source_text": source_text,
    }


def add_file_section(doc, summary, chapter_number: int):
    file_name = summary["file_name"]
    note = FILE_NOTES[file_name]
    chunks: list[Chunk] = summary["chunks"]

    add_heading(doc, f"{chapter_number}. {file_name}", 1)
    add_body(doc, note["purpose"])
    add_body(doc, f"Ruta: {summary['path']}", bold_lead="Ruta:")
    add_body(doc, f"Lineas: {summary['line_count']}", bold_lead="Lineas:")
    add_body(doc, f"Definiciones detectadas: {summary['definition_count']}", bold_lead="Definiciones detectadas:")
    add_body(doc, note["why"], bold_lead=None)

    add_heading(doc, f"{chapter_number}.1 Puntos de lectura importantes", 2)
    for hotspot in note["hotspots"]:
        add_bullet(doc, hotspot)

    add_heading(doc, f"{chapter_number}.2 Mapa de bloques", 2)
    add_table(
        doc,
        ["Rango", "Bloque", "Descripcion"],
        [[f"{chunk.line_start}-{chunk.line_end}", chunk.title, chunk.description] for chunk in chunks],
        [1.1, 2.2, 5.0],
    )

    add_heading(doc, f"{chapter_number}.3 Codigo fuente completo", 2)
    add_body(
        doc,
        "A continuacion se muestra el archivo completo. Antes de cada bloque se incluye una nota de lectura corta para explicar su responsabilidad dentro del flujo general.",
    )

    for chunk_index, chunk in enumerate(chunks, start=1):
        add_heading(
            doc,
            f"{chapter_number}.3.{chunk_index} {chunk.title} [{chunk.line_start}-{chunk.line_end}]",
            3,
        )
        add_body(doc, chunk.description)
        for offset, code_line in enumerate(chunk.lines):
            add_code_line(doc, chunk.line_start + offset, code_line)
        doc.add_paragraph()


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    summaries = [summarize_file(name, path) for name, path in FILES.items()]

    doc = Document()
    doc._monitor_summaries = summaries  # type: ignore[attr-defined]
    set_document_defaults(doc)
    add_cover(doc, summaries)
    add_reading_guide(doc)
    add_file_section(doc, summaries[0], 2)
    add_file_section(doc, summaries[1], 3)
    add_file_section(doc, summaries[2], 4)
    doc.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
