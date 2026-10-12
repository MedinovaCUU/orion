from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/ricardomontanezmiranda/Desktop/Biosystems Project")
MONITOR_DIR = ROOT / "tools" / "ba400-log-monitor"
OUTPUT_DIR = ROOT / "outputs" / "monitor-docs"
OUTPUT_PATH = OUTPUT_DIR / "Documentacion_Tecnica_AX00_Log_Monitor.docx"

FILES = {
    "full-monitor.mjs": MONITOR_DIR / "full-monitor.mjs",
    "monitor.mjs": MONITOR_DIR / "monitor.mjs",
    "consumption-monitor.mjs": MONITOR_DIR / "consumption-monitor.mjs",
}

TITLE_BLUE = RGBColor(0x2E, 0x74, 0xB5)
TITLE_DARK = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "E8EEF5"


SUMMARY_NOTES = {
    "full-monitor.mjs": {
        "purpose": "Supervisor del sistema. Arranca los monitores de errores y consumos, captura sus logs y reinicia procesos caidos.",
        "key_points": [
            "Expande placeholders de configuracion como ${SCRIPT_DIR}, ${CONFIG_DIR} y variables de entorno.",
            "Crea un archivo de log por monitor y prefija cada linea para saber de que proceso viene.",
            "Reinicia automaticamente los hijos despues de restartDelayMs si terminan inesperadamente.",
            "Permite modo --once para pruebas controladas sin lazo infinito.",
        ],
        "core_functions": [
            "normalizeConfig",
            "pipeWithPrefix",
            "MonitorSupervisor",
            "main",
        ],
    },
    "monitor.mjs": {
        "purpose": "Monitor de errores en tiempo real sobre TraceComm*.LAx00. Lee solo el incremento nuevo del archivo, detecta errores activos y publica historico mas estado actual.",
        "key_points": [
            "Autodetecta la carpeta correcta de logs cuando hay varias instalaciones posibles.",
            "Mantiene offset, linea actual, serie detectada y firma del error activo en un estado local persistente.",
            "Da prioridad a ASN sobre SN para capturar el numero de serie real del analizador.",
            "Interpreta mensajes TraceComm, deduplica errores activos y usa error-catalog.json para descripcion y severidad.",
            "Inserta eventos historicos y hace upsert del estado actual por numero de serie en Supabase.",
        ],
        "core_functions": [
            "resolveSourceDirectory",
            "rankErrorDirectoryCandidates",
            "readIncrementalChunk",
            "parseTraceCommLine",
            "buildEvent",
            "buildCurrentErrorStateRow",
            "uploadEvents",
            "uploadCurrentStates",
            "pollOnce",
        ],
    },
    "consumption-monitor.mjs": {
        "purpose": "Monitor de consumos desde LogConsum. Recalcula acumulados mensuales por reactivo, rotor y estado de insumos para que las tablas online sean compactas y reescribibles.",
        "key_points": [
            "Autodetecta la carpeta de LogConsum y escoge la mejor candidata por vigencia del mes actual.",
            "Decide si reprocesar un archivo segun firma local size:mtime, mes actual y resincronizacion periodica.",
            "Agrupa los reactivos por mes + numero_serie + test_name y los rotores por mes + numero_serie.",
            "Actualiza un estado operativo por equipo con PackISE_SN y fechas de electrodos.",
            "Usa upsert para reescribir totales absolutos y evitar duplicados ante reinstalaciones o caidas.",
        ],
        "core_functions": [
            "rankConsumptionDirectoryCandidates",
            "resolveCandidateFiles",
            "shouldReprocessFile",
            "processReagentFile",
            "processRotorFile",
            "upsertRows",
            "pollOnce",
        ],
    },
}

FUNCTION_NOTES = {
    "full-monitor.mjs": {
        "parseArgs": "Lee --config y --once para definir como arrancar el supervisor.",
        "readJson": "Carga el JSON de configuracion del supervisor.",
        "timestamp": "Genera marcas de tiempo ISO para todos los logs.",
        "log": "Estandariza la salida del supervisor.",
        "sleep": "Espera asincrona usada antes de reiniciar monitores.",
        "expandPlaceholders": "Resuelve placeholders de ruta y variables de entorno.",
        "deepExpandPlaceholders": "Aplica la expansion de placeholders sobre estructuras anidadas.",
        "normalizeConfig": "Convierte el JSON de entrada en una configuracion lista para ejecutar.",
        "ensureLogDir": "Crea la carpeta de logs si aun no existe.",
        "pipeWithPrefix": "Recibe stdout/stderr de cada hijo, los separa por linea y los reescribe con prefijo.",
        "MonitorSupervisor": "Clase central que lanza, vigila, reinicia y detiene los procesos hijos.",
        "main": "Entrada principal: carga config, crea el supervisor y maneja apagado limpio.",
    },
    "monitor.mjs": {
        "parseArgs": "Lee los argumentos de ejecucion del monitor de errores.",
        "readJson": "Carga archivos JSON como configuracion y catalogo.",
        "writeJson": "Persiste el estado local del monitor.",
        "timestamp": "Genera timestamps de log.",
        "log": "Escribe mensajes del monitor a consola/log.",
        "toLocalDateKey": "Convierte una fecha al formato YYYYMMDD usado al priorizar TraceComm del dia.",
        "sleep": "Pausa corta usada en reintentos de lectura bloqueada.",
        "expandPlaceholders": "Expande placeholders de configuracion.",
        "deepExpandPlaceholders": "Expande placeholders dentro de objetos y arreglos.",
        "getValueByPath": "Extrae valores anidados para el columnMap del upload.",
        "removeUndefinedProperties": "Limpia campos undefined antes de enviar datos.",
        "inferEquipmentModelFromSerial": "Infiere BA400, BA200, A15 o A25 segun prefijo de serie.",
        "extractPreferredTraceCommSerial": "Busca ASN y SN en una linea o bloque; ASN tiene prioridad.",
        "normalizeCatalogErrorType": "Normaliza fatal/warning/ok del catalogo.",
        "resolveMessageTypeFromCatalog": "Escoge la severidad mas alta cuando hay varios codigos a la vez.",
        "fileExists": "Comprueba si el archivo de estado ya existe.",
        "loadState": "Carga el estado local del monitor de errores.",
        "resolveSourceDirectory": "Resuelve o autodetecta la carpeta correcta de logs.",
        "extractTraceDateKeyFromName": "Extrae la fecha embebida en TraceCommYYYYMMDD.*",
        "rankErrorDirectoryCandidates": "Ordena carpetas candidatas privilegiando el dia actual y la actividad reciente.",
        "resolveActiveFile": "Escoge el archivo TraceComm mas reciente dentro de la carpeta activa.",
        "getOrCreateFileState": "Devuelve el estado asociado a un archivo concreto y lo inicializa si hace falta.",
        "readIncrementalChunk": "Lee solo el rango de bytes nuevo desde el ultimo offset.",
        "readTextSlice": "Convierte un rango de bytes en texto con el encoding configurado.",
        "extractTraceCommIdentity": "Puente para extraer identidad del analizador desde texto TraceComm.",
        "bootstrapTraceCommIdentity": "Escanea cabeza y cola del archivo para aprender la serie aun si no hay errores nuevos.",
        "parseErrorCodes": "Extrae codigos con el regex configurado en modo simple.",
        "parseTraceCommLine": "Interpreta la linea TraceComm, detecta errores activos, limpiezas y deduplicacion.",
        "buildMatchedDescriptions": "Cruza cada codigo con el catalogo local.",
        "resolveCurrentStateStatus": "Resuelve el estado actual del equipo: ok, warning o fatal.",
        "buildEvent": "Construye el payload historico para monitoreo_errores_equipos.",
        "buildCurrentErrorStateRow": "Construye la fila consolidada para estado_errores_equipo_actual.",
        "mapEventToRow": "Aplica columnMap para adaptar el evento interno al esquema SQL.",
        "uploadBatch": "Inserta historicos por lotes en Supabase.",
        "upsertBatch": "Hace upsert sobre tablas de estado actual.",
        "uploadEvents": "Parte eventos en lotes y los inserta.",
        "uploadCurrentStates": "Deduplica estados por serie y hace upsert.",
        "pollOnce": "Ejecuta un ciclo completo de deteccion, parseo, estado y carga.",
        "normalizeConfig": "Normaliza toda la configuracion del monitor de errores.",
        "main": "Bucle principal con manejo de errores y frecuencia de sondeo.",
    },
    "consumption-monitor.mjs": {
        "parseArgs": "Lee --config y --once para el monitor de consumos.",
        "readJson": "Carga configuraciones y estado local.",
        "writeJson": "Guarda el estado local de consumos.",
        "fileExists": "Comprueba si existe el archivo de estado.",
        "timestamp": "Genera timestamps para logging.",
        "log": "Registra actividad del monitor.",
        "sleep": "Pausa corta usada en reintentos de lectura.",
        "expandPlaceholders": "Expande placeholders de configuracion.",
        "deepExpandPlaceholders": "Expansion profunda de placeholders.",
        "inferEquipmentModelFromSerial": "Deduce el modelo a partir de la serie.",
        "normalizeText": "Limpia nulos y espacios de texto crudo.",
        "normalizeCsvCell": "Limpia una celda CSV, incluyendo tabs residuales.",
        "parseNumber": "Convierte texto numerico de CSV a numero.",
        "parseBoolean": "Normaliza booleanos como TRUE/1/YES.",
        "parseLocalTimestamp": "Convierte la columna Date del CSV en fecha valida.",
        "sha1": "Genera llaves estables summary_key y pipeline_signature.",
        "buildPipelineSignature": "Firma del pipeline para forzar resincronizacion si cambia el contrato de salida.",
        "parseCsvLine": "Parser CSV simple con manejo de comillas.",
        "mapCsvRow": "Mapea encabezados a valores de una fila.",
        "normalizeSampleClassKey": "Normaliza SampleClass a BLANK/CALIB/CTRL/PATIENT/OTHER.",
        "toLocalMonthKey": "Genera YYYYMM para el mes actual.",
        "readTextFileWithRetries": "Lee archivos de consumo con reintentos ante EBUSY/EACCES/EPERM.",
        "parseConsumptionFileName": "Extrae serie, tipo y mes desde el nombre del archivo.",
        "extractConsumptionMonthKey": "Saca el mes YYYYMM del nombre del archivo.",
        "extractConsumptionKind": "Distingue archivos de reactivo y rotor.",
        "rankConsumptionDirectoryCandidates": "Ordena carpetas segun vigencia del mes actual y actividad reciente.",
        "resolveSourceDirectory": "Resuelve o autodetecta la carpeta correcta de LogConsum.",
        "resolveCandidateFiles": "Lista todos los archivos de consumo validos dentro de la carpeta activa.",
        "loadState": "Carga el estado local del monitor de consumos.",
        "getOrCreateFileState": "Inicializa o devuelve el estado de un archivo mensual.",
        "shouldReprocessFile": "Decide si debe releerse un archivo por cambio, mes actual o resincronizacion.",
        "ensureEquipmentState": "Inicializa el acumulador de estado operativo por equipo.",
        "updateEquipmentStateFromRow": "Actualiza PackISE_SN y electrodos con la fila mas reciente.",
        "buildEquipmentStateRow": "Convierte el estado operativo a una fila lista para Supabase.",
        "createReagentSummary": "Inicializa un acumulador mensual por prueba.",
        "updateReagentSummary": "Suma conteos y volumenes sobre cada pipeteo.",
        "finalizeReagentSummary": "Cierra el acumulado de reactivo con redondeos y metadatos finales.",
        "createRotorSummary": "Inicializa el acumulado mensual de cambios de rotor.",
        "updateRotorSummary": "Incrementa conteo de cambios y conserva primer/ultimo evento.",
        "finalizeRotorSummary": "Cierra la fila de rotor para carga.",
        "resolveEffectiveIdentity": "Escoge la serie efectiva desde AnalyzerSN, config o nombre de archivo.",
        "mergeEquipmentStateMaps": "Une estados parciales y conserva el evento mas reciente por equipo.",
        "processReagentFile": "Lee un CSV de reactivos y produce acumulados mensuales y estado operativo.",
        "processRotorFile": "Lee un CSV de rotores y produce acumulados mensuales por equipo.",
        "processFile": "Despacha al parser correcto segun el tipo de archivo.",
        "chunkRows": "Parte arreglos grandes en lotes de upsert.",
        "upsertRows": "Hace upsert de filas en Supabase.",
        "normalizeConfig": "Normaliza la configuracion del monitor de consumos.",
        "pollOnce": "Ejecuta un ciclo completo de reprocesamiento, agregacion y sincronizacion.",
        "main": "Bucle principal con la frecuencia horaria configurada.",
    },
}


def set_run_font(run, *, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.1):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_text(cell, text, *, bold=False, color=None, font_size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    set_paragraph_spacing(paragraph, before=0, after=0, line=1.1)
    run = paragraph.add_run(text)
    set_run_font(run, size=font_size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_fixed_layout(table, widths_in_inches, *, indent_dxa=120, header_fill=LIGHT_FILL):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    total_dxa = sum(int(width * 1440) for width in widths_in_inches)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    width_dxa_values = [int(width * 1440) for width in widths_in_inches]
    for width_dxa in width_dxa_values:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width_dxa))
        tbl_grid.append(grid_col)

    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            width_dxa = width_dxa_values[col_idx]
            cell.width = Inches(widths_in_inches[col_idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")

            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, margin in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = tc_mar.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    tc_mar.append(node)
                node.set(qn("w:w"), str(margin))
                node.set(qn("w:type"), "dxa")

            if row_idx == 0 and header_fill:
                shade_cell(cell, header_fill)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(8)
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)
    else:
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=6, line=1.1)
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        set_run_font(run, size=11, bold=True)
        run2 = paragraph.add_run(text[len(bold_lead):])
        set_run_font(run2, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_table(doc, headers, rows, widths, *, header_fill=LIGHT_FILL, style="Table Grid", indent_dxa=120):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(hdr[idx], header, bold=True, color=TITLE_DARK, font_size=10)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), font_size=10)
    set_table_fixed_layout(table, widths, indent_dxa=indent_dxa, header_fill=header_fill)
    doc.add_paragraph()
    return table


def add_metadata_row(doc, label, value):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=2, line=1.05)
    run1 = paragraph.add_run(f"{label}: ")
    set_run_font(run1, size=11, bold=True)
    run2 = paragraph.add_run(value)
    set_run_font(run2, size=11)


def insert_page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def set_document_defaults(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for level, size, color in (
        (1, 16, TITLE_BLUE),
        (2, 13, TITLE_BLUE),
        (3, 12, TITLE_DARK),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(header_para, before=0, after=0, line=1.0)
    header_run = header_para.add_run("AX00 Log Monitor | Documentacion tecnica del codigo fuente")
    set_run_font(header_run, size=9, color=GRAY)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer_para, before=0, after=0, line=1.0)
    footer_run = footer_para.add_run("Generado el 20/08/2026 a partir del codigo fuente del proyecto")
    set_run_font(footer_run, size=8.5, color=GRAY)


def extract_definitions(source_path: Path):
    text = source_path.read_text(encoding="utf8")
    definitions = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        match = re.match(r"^(function|async function|class)\s+([A-Za-z0-9_]+)", line)
        if match:
            definitions.append(
                {
                    "kind": match.group(1),
                    "name": match.group(2),
                    "line": line_number,
                }
            )
    return text, definitions


def summarize_file(file_name: str, source_path: Path):
    text, definitions = extract_definitions(source_path)
    lines = text.splitlines()
    return {
        "file_name": file_name,
        "path": str(source_path),
        "line_count": len(lines),
        "definitions": definitions,
        "function_count": len(definitions),
        "class_count": len([d for d in definitions if d["kind"] == "class"]),
    }


def add_cover(doc, summaries):
    add_body(doc, "Documentacion tecnica del monitor AX00", bold_lead=None)
    cover_title = doc.paragraphs[-1]
    cover_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cover_title.paragraph_format.space_before = Pt(18)
    cover_title.paragraph_format.space_after = Pt(4)
    for run in cover_title.runs:
        set_run_font(run, size=23, color=RGBColor(0, 0, 0), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run(
        "Explicacion detallada de full-monitor.mjs, monitor.mjs y consumption-monitor.mjs"
    )
    set_run_font(subtitle_run, size=14, color=GRAY)

    add_metadata_row(doc, "Proyecto", "AX00 Log Monitor / Monitoreo nacional de equipos BA200, BA400, A15 y A25")
    add_metadata_row(doc, "Fecha", "20 de agosto de 2026")
    add_metadata_row(doc, "Objetivo", "Dejar documentada la arquitectura, el flujo y las funciones principales del codigo fuente")
    add_metadata_row(doc, "Archivos analizados", ", ".join(summary["file_name"] for summary in summaries))

    doc.add_paragraph()
    add_body(
        doc,
        "Este documento describe como se coordina el monitoreo completo, como se detectan y sincronizan los errores activos desde TraceComm y como se recalculan los consumos mensuales desde LogConsum. La explicacion esta orientada al codigo fuente real del proyecto para que sirva como base de mantenimiento, soporte e incorporacion de nuevas funciones.",
    )


def add_architecture_section(doc, summaries):
    add_heading(doc, "1. Arquitectura general del paquete", 1)
    add_body(
        doc,
        "El sistema esta compuesto por tres piezas principales. full-monitor.mjs actua como supervisor; monitor.mjs se encarga de errores en tiempo real desde TraceComm; y consumption-monitor.mjs genera acumulados mensuales desde LogConsum. Todas las piezas comparten la misma filosofia: detectar automaticamente la fuente correcta, mantener estado local, registrar logs legibles y publicar datos consistentes en Supabase.",
    )

    add_heading(doc, "1.1 Responsabilidad de cada archivo", 2)
    rows = []
    for summary in summaries:
        note = SUMMARY_NOTES[summary["file_name"]]
        rows.append(
            [
                summary["file_name"],
                summary["line_count"],
                note["purpose"],
            ]
        )
    add_table(
        doc,
        ["Archivo", "Lineas", "Responsabilidad principal"],
        rows,
        [2.0, 0.7, 3.8],
        header_fill=LIGHT_BLUE_FILL,
    )

    add_heading(doc, "1.2 Flujo global de ejecucion", 2)
    for step in (
        "El sistema arranca desde run-full-monitor.cmd o desde la tarea programada de Windows.",
        "full-monitor.mjs lee config.windows.completo.json y lanza dos procesos hijos: errores y consumos.",
        "monitor.mjs localiza el TraceComm activo, lee solo el delta nuevo y detecta cambios de estado de error.",
        "consumption-monitor.mjs localiza LogConsum, decide que archivos mensual(es) requieren reproceso y recalcula los totales absolutos.",
        "Ambos procesos guardan estado local en la carpeta state y escriben logs legibles en la carpeta logs.",
        "Si uno de los monitores cae, full-monitor.mjs lo reinicia automaticamente sin afectar al otro.",
    ):
        add_number(doc, step)

    add_heading(doc, "1.3 Dependencias funcionales", 2)
    add_bullet(doc, "source-discovery.mjs: utilitario compartido para autodeteccion de carpetas validas.")
    add_bullet(doc, "error-catalog.json: catalogo local de codigos, descripcion, seccion y severidad.")
    add_bullet(doc, "config.windows.*.json: describen origenes, patrones, frecuencia de sondeo, tablas destino y estado local.")
    add_bullet(doc, "Supabase REST API: destino de historicos, estados actuales y acumulados de consumos.")


def add_file_section(doc, summary):
    note = SUMMARY_NOTES[summary["file_name"]]
    add_heading(doc, f"2. {summary['file_name']}" if summary["file_name"] == "full-monitor.mjs" else "", 1)


def add_full_monitor_section(doc, summary):
    add_heading(doc, "2. full-monitor.mjs", 1)
    add_body(doc, SUMMARY_NOTES[summary["file_name"]]["purpose"])
    add_heading(doc, "2.1 Que problema resuelve", 2)
    add_body(
        doc,
        "El supervisor existe para que el despliegue sea estable en Windows. En lugar de depender de una sola consola o de un proceso monolitico, desacopla errores y consumos en procesos independientes. Esto permite reiniciar uno sin tumbar el otro y separar claramente los logs de cada monitor.",
    )

    add_heading(doc, "2.2 Flujo interno", 2)
    for item in SUMMARY_NOTES[summary["file_name"]]["key_points"]:
        add_bullet(doc, item)

    add_heading(doc, "2.3 Clase central: MonitorSupervisor", 2)
    add_body(
        doc,
        "La clase MonitorSupervisor encapsula el ciclo de vida de los monitores hijos. Su mapa interno children guarda el proceso, su configuracion y el stream de log abierto por nombre de monitor. startAll() crea la carpeta de logs y llama startMonitor() por cada monitor habilitado; stopAll() manda SIGTERM a todos los hijos; waitUntilDone() y resolveIfDone() coordinan un cierre limpio.",
    )
    add_body(
        doc,
        "startMonitor() construye el comando node <script> --config <config>, abre un log dedicado, enlaza stdout y stderr mediante pipeWithPrefix() y registra un listener de salida. Si el proceso termina y el supervisor no esta apagandose, espera restartDelayMs y lo vuelve a lanzar.",
    )

    add_heading(doc, "2.4 Funciones principales", 2)
    rows = []
    for definition in summary["definitions"]:
        rows.append(
            [
                str(definition["line"]),
                definition["name"],
                FUNCTION_NOTES[summary["file_name"]].get(definition["name"], "Helper interno del supervisor."),
            ]
        )
    add_table(doc, ["Linea", "Funcion o clase", "Responsabilidad"], rows, [0.7, 1.8, 4.0])

    add_heading(doc, "2.5 Puntos de mantenimiento", 2)
    add_bullet(doc, "Si se agrega un nuevo monitor al paquete, normalmente basta con declararlo en config.windows.completo.json.")
    add_bullet(doc, "restartDelayMs define cuanto espera el supervisor antes de relanzar un hijo caido.")
    add_bullet(doc, "Los placeholders de rutas permiten mover la carpeta completa sin hardcodear paths absolutos.")


def add_monitor_section(doc, summary):
    add_heading(doc, "3. monitor.mjs", 1)
    add_body(doc, SUMMARY_NOTES[summary["file_name"]]["purpose"])

    add_heading(doc, "3.1 Objetivo operativo", 2)
    add_body(
        doc,
        "Su objetivo no es leer un log completo cada vez, sino comportarse como un seguidor incremental del TraceComm activo. El monitor parte del supuesto de que el archivo puede estar abierto por otro programa, crecer continuamente, cambiar de nombre cada dia y contener datos de telemetria mezclados con errores y estados internos.",
    )

    add_heading(doc, "3.2 Resolucion del origen", 2)
    add_body(
        doc,
        "resolveSourceDirectory() usa source-discovery.mjs para evaluar rutas preferidas y raices de busqueda. Luego rankErrorDirectoryCandidates() aplica la regla especifica del monitor: priorizar carpetas que tengan un TraceComm con la fecha local del dia, y solo despues comparar mtime, cantidad de archivos validos y nombre de carpeta.",
    )
    add_body(
        doc,
        "Una vez elegida la carpeta, resolveActiveFile() selecciona el archivo mas reciente cuyo nombre cumpla fileNamePattern. Esto permite soportar rotacion diaria de TraceComm sin exigir una ruta exacta al archivo del dia.",
    )

    add_heading(doc, "3.3 Estado local y lectura incremental", 2)
    add_body(
        doc,
        "Cada archivo observado tiene su propio bloque de estado: offset, line_number, pending_fragment, serie detectada, firma del error activo y marca initialized. Ese estado se guarda en JSON y permite reiniciar el proceso sin perder continuidad.",
    )
    add_body(
        doc,
        "readIncrementalChunk() lee solo el rango [offset anterior, tamano actual] y tolera bloqueos transitorios del archivo con reintentos. Si la ultima lectura termina en una linea incompleta, el texto restante se conserva en pending_fragment y se completa en el siguiente ciclo antes de separar lineas.",
    )

    add_heading(doc, "3.4 Deteccion de serie y modelo", 2)
    add_body(
        doc,
        "El monitor necesita identificar correctamente el equipo antes de publicar. extractPreferredTraceCommSerial() busca ASN y SN; si ambos existen, ASN tiene prioridad porque representa la serie real del analizador y evita falsos cambios introducidos por lineas como ANSADJ. bootstrapTraceCommIdentity() escanea tanto la cabeza como la cola del archivo para aprender la serie aunque el ciclo actual no traiga errores nuevos.",
    )
    add_body(
        doc,
        "inferEquipmentModelFromSerial() mapea los prefijos conocidos a BA400, BA200, A15 y A25. Eso evita hardcodear el modelo por instalacion.",
    )

    add_heading(doc, "3.5 Parser TraceComm y maquina de estados", 2)
    add_body(
        doc,
        "parseTraceCommLine() es la funcion mas importante del archivo. Analiza una sola linea y produce tres cosas: errorCodes, metadata y stateTransition. Primero identifica analyzer_id y message_type; luego recoge todos los E:<codigo>; despues calcula si la linea representa un borrado de errores (por ejemplo E:0 o ANSERR;N:0) o una permanencia/activacion de errores reales.",
    )
    add_body(
        doc,
        "La funcion maneja dos firmas distintas. tracecomm_active_error_signature sirve para deduplicar el mismo error activo repetido en varias lineas consecutivas. tracecomm_last_reported_state_signature sirve para no repetir la misma transicion de estado actual. Con eso el monitor puede evitar spam de eventos sin dejar de publicar el cambio a limpio o el cambio a un nuevo error.",
    )

    add_heading(doc, "3.6 Construccion de datos de salida", 2)
    add_bullet(doc, "buildEvent(): crea el registro historico para monitoreo_errores_equipos.")
    add_bullet(doc, "buildCurrentErrorStateRow(): crea el estado consolidado para estado_errores_equipo_actual.")
    add_bullet(doc, "mapEventToRow(): aplica el columnMap definido en la configuracion y permite adaptar el payload al esquema SQL sin tocar el parser.")

    add_body(
        doc,
        "El historico conserva detalle de trazabilidad como source_file, line_number, byte offsets, raw_line, line_hash y payload completo. El estado actual, en cambio, privilegia el resumen operativo del equipo: estado_actual, codigo_estado, descripcion_estado, errores_activos, error_principal_codigo y timestamps de actividad.",
    )

    add_heading(doc, "3.7 Carga a Supabase", 2)
    add_body(
        doc,
        "uploadBatch() inserta los eventos historicos por lotes. upsertBatch() hace merge sobre la tabla de estado actual usando on_conflict. uploadCurrentStates() deduplica previamente por numero_serie para que solo sobreviva el ultimo estado calculado en el ciclo.",
    )

    add_heading(doc, "3.8 Secuencia completa de pollOnce()", 2)
    for step in (
        "Resuelve la carpeta fuente y el archivo activo.",
        "Carga o crea fileState para ese archivo.",
        "Si es un archivo nuevo y startAtEndOnFirstSeen esta activo, salta el historico previo.",
        "Si el archivo se truncó o rotó, reinicia offset y linea.",
        "Lee solo el bloque nuevo desde el ultimo offset.",
        "Separa lineas, recompone fragmentos pendientes y procesa una a una.",
        "Genera eventos historicos y transiciones de estado actual.",
        "Actualiza el estado local en disco.",
        "Sube primero eventos historicos y luego el estado actual.",
    ):
        add_number(doc, step)

    add_heading(doc, "3.9 Inventario de funciones", 2)
    rows = []
    for definition in summary["definitions"]:
        rows.append(
            [
                str(definition["line"]),
                definition["name"],
                FUNCTION_NOTES[summary["file_name"]].get(definition["name"], "Helper interno del monitor."),
            ]
        )
    add_table(doc, ["Linea", "Funcion", "Descripcion"], rows, [0.7, 1.9, 3.9], header_fill=LIGHT_BLUE_FILL)


def add_consumption_section(doc, summary):
    add_heading(doc, "4. consumption-monitor.mjs", 1)
    add_body(doc, SUMMARY_NOTES[summary["file_name"]]["purpose"])

    add_heading(doc, "4.1 Diferencia conceptual frente al monitor de errores", 2)
    add_body(
        doc,
        "Este monitor no sigue bytes nuevos de un archivo en crecimiento. En cambio, recorre CSV mensuales que ya representan hechos completos de consumo. Por eso la estrategia es recalcular totales absolutos y hacer upsert de resumenes, en vez de subir cada pipeteo como un evento online independiente.",
    )

    add_heading(doc, "4.2 Seleccion del directorio y de los archivos", 2)
    add_body(
        doc,
        "resolveSourceDirectory() y rankConsumptionDirectoryCandidates() aplican una heuristica distinta a la de errores: favorecen carpetas que tengan archivos de reactivos o rotor del mes actual. resolveCandidateFiles() recorre la carpeta activa y construye una lista de candidatos con tipo, mes, serie, tamano, mtime y firma local.",
    )

    add_heading(doc, "4.3 Estado local y criterio de reproceso", 2)
    add_body(
        doc,
        "El estado local guarda por archivo la ultima firma observada y el ultimo synced_at. shouldReprocessFile() devuelve true si cambió la firma del archivo, si el archivo corresponde al mes actual o si vencio la resincronizacion periodica. Esto hace que los meses en curso se recalculen siempre y los meses cerrados solo cuando tiene sentido.",
    )
    add_body(
        doc,
        "buildPipelineSignature() protege el contrato de salida. Si cambia la definicion del pipeline, las tablas destino o la estrategia de agregacion, pollOnce() invalida el estado local y fuerza una resincronizacion completa de consumos.",
    )

    add_heading(doc, "4.4 Parser de reactivos", 2)
    add_body(
        doc,
        "processReagentFile() interpreta la primera fila como encabezado y luego procesa cada pipeteo. A partir de cada fila obtiene una identidad efectiva, normaliza TestName y SampleClass, calcula el timestamp del evento y actualiza un summary mensual por llave mes + numero_serie + test_name.",
    )
    add_bullet(doc, "pipetting_count: numero total de pipeteos para esa prueba en el mes.")
    add_bullet(doc, "vr1_total_ul y vr2_total_ul: volumen total de reactivo consumido.")
    add_bullet(doc, "sample_volume_total_ul: volumen total de muestra aspirado.")
    add_bullet(doc, "blank_count, calib_count, ctrl_count, patient_count: distribucion por clase de muestra.")
    add_bullet(doc, "factory_test_count y non_factory_test_count: separacion entre pruebas de fabrica y no fabrica.")

    add_heading(doc, "4.5 Estado operativo por equipo", 2)
    add_body(
        doc,
        "Mientras procesa reactivos, el monitor mantiene un mapa equipmentStates por numero_serie. updateEquipmentStateFromRow() aprovecha el mismo CSV para capturar el estado reciente de PackISE_SN y las fechas de electrodos Ref, Na, K, Cl y Li. El resultado final se materializa con buildEquipmentStateRow() y se hace upsert por serie.",
    )

    add_heading(doc, "4.6 Parser de rotores", 2)
    add_body(
        doc,
        "processRotorFile() es mas simple: cada fila valida representa un cambio de rotor. El monitor agrupa por mes + numero_serie y acumula rotor_change_count, first_change_at, last_change_at y change_timestamps. Esa tabla sirve como resumen mensual y evidencia auditable.",
    )

    add_heading(doc, "4.7 Secuencia completa de pollOnce()", 2)
    for step in (
        "Calcula pipeline_signature y reinicia estado si el contrato de salida cambió.",
        "Resuelve la carpeta LogConsum activa.",
        "Enumera los archivos candidatos de reactivos y rotores.",
        "Decide cuales requieren reproceso con shouldReprocessFile().",
        "Procesa cada archivo y llena tres mapas: reactivos, rotores y estado de equipo.",
        "Hace upsert por lotes en las tres tablas destino.",
        "Solo si la carga termina bien actualiza last_signature y last_synced_at de cada archivo procesado.",
    ):
        add_number(doc, step)

    add_heading(doc, "4.8 Inventario de funciones", 2)
    rows = []
    for definition in summary["definitions"]:
        rows.append(
            [
                str(definition["line"]),
                definition["name"],
                FUNCTION_NOTES[summary["file_name"]].get(definition["name"], "Helper interno del monitor de consumos."),
            ]
        )
    add_table(doc, ["Linea", "Funcion", "Descripcion"], rows, [0.7, 1.9, 3.9], header_fill=LIGHT_BLUE_FILL)


def add_cross_section(doc):
    add_heading(doc, "5. Relacion entre las piezas", 1)
    add_body(
        doc,
        "Las tres piezas forman una arquitectura de observabilidad local con dos contratos de datos distintos. full-monitor.mjs se preocupa por disponibilidad del servicio; monitor.mjs produce eventos historicos y estado actual de errores; consumption-monitor.mjs produce acumulados mensuales y estado operativo de insumos. Ninguno de los monitores depende internamente del otro, pero ambos se apoyan en convenciones compartidas: autodeteccion de fuentes, estado local en JSON y sincronizacion REST contra Supabase.",
    )

    add_heading(doc, "5.1 Diferencias clave entre errores y consumos", 2)
    add_table(
        doc,
        ["Tema", "monitor.mjs", "consumption-monitor.mjs"],
        [
            ["Fuente principal", "TraceComm*.LAx00 en crecimiento", "CSV mensuales de LogConsum"],
            ["Unidad logica", "Evento de error y estado actual", "Acumulado mensual por prueba o rotor"],
            ["Lectura", "Incremental por bytes y offsets", "Relectura completa por archivo cuando hace falta"],
            ["Persistencia local", "offset, fragmento, firmas de error", "firma por archivo y ultimo synced_at"],
            ["Destino SQL", "Historico + estado actual", "Reactivos, rotores y estado operativo"],
            ["Estrategia de duplicados", "Deduplicacion de error activo", "Upsert de totales absolutos por llave estable"],
        ],
        [1.5, 2.45, 2.55],
        header_fill=LIGHT_BLUE_FILL,
    )

    add_heading(doc, "5.2 Logs y diagnostico", 2)
    add_bullet(doc, "logs/errores.log: salida especifica del monitor de errores.")
    add_bullet(doc, "logs/consumos.log: salida especifica del monitor de consumos.")
    add_bullet(doc, "logs/full-monitor-daemon.log: log del wrapper que se deja corriendo al inicio de Windows.")
    add_bullet(doc, "state/monitor-state-*.json: permite depurar desde donde se esta leyendo.")
    add_bullet(doc, "state/consumption-monitor-state.json: permite verificar por que un archivo no se reproceso.")

    add_heading(doc, "5.3 Riesgos y zonas sensibles del codigo", 2)
    add_bullet(doc, "Cambiar los patrones de archivos o el ranking de carpetas afecta directamente la instalacion nacional con rutas no estandar.")
    add_bullet(doc, "Cambiar columnMap o las llaves de upsert modifica el contrato con Supabase y puede requerir resincronizacion.")
    add_bullet(doc, "Cambiar la logica de ASN/SN puede introducir series incorrectas en equipos con mensajes ANSADJ.")
    add_bullet(doc, "Cambiar la agregacion de consumos exige revisar pipeline_signature para forzar reproceso coherente.")


def add_appendix(doc, summaries):
    add_heading(doc, "6. Anexo: resumen ejecutivo por archivo", 1)
    for summary in summaries:
        note = SUMMARY_NOTES[summary["file_name"]]
        add_heading(doc, summary["file_name"], 2)
        add_body(doc, f"Ruta: {summary['path']}", bold_lead="Ruta:")
        add_body(doc, f"Lineas: {summary['line_count']}", bold_lead="Lineas:")
        add_body(doc, f"Funciones/clases detectadas: {summary['function_count']}", bold_lead="Funciones/clases detectadas:")
        add_body(doc, f"Responsabilidad: {note['purpose']}", bold_lead="Responsabilidad:")
        add_body(
            doc,
            "Funciones nucleo: " + ", ".join(note["core_functions"]),
            bold_lead="Funciones nucleo:",
        )


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    summaries = [summarize_file(name, path) for name, path in FILES.items()]

    doc = Document()
    set_document_defaults(doc)
    add_cover(doc, summaries)
    insert_page_break(doc)
    add_architecture_section(doc, summaries)
    add_full_monitor_section(doc, summaries[0])
    add_monitor_section(doc, summaries[1])
    add_consumption_section(doc, summaries[2])
    add_cross_section(doc)
    add_appendix(doc, summaries)
    doc.save(str(OUTPUT_PATH))
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
